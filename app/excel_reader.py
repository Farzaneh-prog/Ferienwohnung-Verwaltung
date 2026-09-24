"""
Reads reservation rows out of the real GästeListe_*.xlsx workbooks.

Phase 1 scope: read + display only. The core guest/stay columns (A-R area)
are mapped by header name (robust to the minor column differences between
the Karlstraße and Eisenach files). The full financial formula columns
(S-BE) are Phase 3 work per the architecture doc's build order — not
parsed here.

Column headers are matched by exact German text (stripped of surrounding
whitespace) against the real headers confirmed in both GästeListe files,
so this does not depend on fixed column letters.
"""
import datetime
import openpyxl

from .config import QUARTER_SHEETS, FUTURE_YEAR_SHEET, NON_DATA_SHEETS, resolve_property_path
from .known_codes import get_known_codes

# canonical field name -> exact header text in the workbook
FIELD_HEADERS = {
    "guest_name": "Gast Name",
    "platform": "Buchung Platform",
    "checkin": "Anreise Datum",
    "checkout": "Abreise Datum",
    "checkout_day": "Abreise Tag",
    "nights": "Nächte Anzahl",
    "adults": "Erwachsene",
    "children": "Kinder unter 18",
    "tourist_tax": "Übernachtungssteuer",
    "paid": "Gezahlt",
    "cleaning_cost": "Putzarbeit",
    "cleaning_fee_charged": "Endreinigung",
    "earned": "Verdient",
    "foreign_company": "Ausländische Firma mit steuer in Heimat",
    "invoice_number": "Rechnungsnummer",
    "confirmation_code": "Bestätigungs-Code",
    "transferred": "bereits überweist",
    "earned_per_night": "Verdient pro nacht (für mich)",
    "guest_email": "Email von Kunde",
    "cancelled": "Storniert",
    "platform_commission": "Payment Charge von Booking",
    "nettobetrag": "Nettobetrag",
    "vat_rate": "Umsatzsteuersatz",
}


def _build_header_map(header_row):
    """header text (stripped) -> 0-based column index."""
    header_map = {}
    for idx, cell in enumerate(header_row):
        if cell is None:
            continue
        text = str(cell).strip()
        if text:
            header_map[text] = idx
    return header_map


def _row_to_reservation(row, header_map, property_key, sheet_name):
    def get(field_key):
        header_text = FIELD_HEADERS[field_key]
        idx = header_map.get(header_text)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    guest_name = get("guest_name")
    if not guest_name:
        return None
    # A real reservation always has both dates. The Muster sheet has a
    # tax-summary block (Umsatzsteuerrechnung) sitting below the real
    # reservations, with row labels like "Zu Zahlende Betrag an dem
    # Behörden" in the same column as Gast Name and no dates — without
    # this check, those got misread as fake reservations with no
    # checkin/checkout, inflating /reservations and find_incomplete_
    # reservations (found 2026-09-24, e.g. "155 Reservierungen" for
    # Karlstraße when the real count is far lower).
    if not get("checkin") or not get("checkout"):
        return None

    reservation = {"property": property_key, "sheet": sheet_name}
    for field_key in FIELD_HEADERS:
        reservation[field_key] = get(field_key)
    return reservation


def _sort_key(reservation):
    checkin = reservation.get("checkin")
    if isinstance(checkin, (datetime.date, datetime.datetime)):
        return checkin
    if isinstance(checkin, str):
        for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
            try:
                return datetime.datetime.strptime(checkin, fmt)
            except ValueError:
                continue
    # Unparseable/missing dates sort last, but don't crash the page.
    return datetime.datetime.max


def find_incomplete_reservations(property_key: str) -> list:
    """Reservations imported via CSV/XLS that still need a manual pass
    (placeholder guest name — see app/import_parser.py)."""
    from .import_parser import is_placeholder_name  # local import: avoids a cycle

    return [r for r in load_reservations(property_key) if is_placeholder_name(r.get("guest_name"))]


def get_existing_confirmation_codes(property_key: str) -> set:
    """All confirmation codes already present for a property, normalized to
    str, for de-duplicating imports (a CSV export can overlap what's already
    in the sheet — e.g. Airbnb's exports can't be filtered by booking date).

    Also includes app/known_codes.py's persistent record — codes this app
    has ever added or cancelled, kept even after Farzaneh manually deletes
    a cancelled row from GästeListe (2026-09-23 decision: she deletes
    cancellations outright, rather than leaving Storniert=ja in place, so
    the un-SUMIFS'd yearly total formulas don't double-count them; without
    this second record, a re-uploaded old export would look like a brand
    new reservation and get re-added — see git history 2026-09-22/23)."""
    codes = set()
    for r in load_reservations(property_key):
        code = r.get("confirmation_code")
        if code not in (None, ""):
            codes.add(str(code).strip())
    codes |= get_known_codes(property_key)
    return codes


def load_reservations(property_key: str) -> list:
    """Load + merge all quarter sheets for one property, sorted by Anreise Datum."""
    path = resolve_property_path(property_key)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        reservations = []
        for sheet_name in QUARTER_SHEETS + [FUTURE_YEAR_SHEET]:
            if sheet_name not in workbook.sheetnames or sheet_name in NON_DATA_SHEETS:
                continue
            sheet = workbook[sheet_name]
            rows = sheet.iter_rows(values_only=True)
            try:
                header_row = next(rows)
            except StopIteration:
                continue
            header_map = _build_header_map(header_row)
            for row in rows:
                reservation = _row_to_reservation(row, header_map, property_key, sheet_name)
                if reservation:
                    reservations.append(reservation)
        reservations.sort(key=_sort_key)
        return reservations
    finally:
        workbook.close()
